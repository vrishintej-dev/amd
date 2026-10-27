"""Corpus walking, parsing and chunking.

Every file is handled inside its own try/except: one unreadable, encrypted,
corrupt or unknown file must never stop the walk. Nothing here needs the GPU;
images (and scanned PDF pages) are handed to an `ocr` callback supplied by the
server, which runs the vision model.
"""

from __future__ import annotations

import csv
import io
import os
import re
import time
from dataclasses import dataclass, field, asdict
from pathlib import Path
from typing import Callable, Optional

TEXT_EXT = {
    ".txt", ".log", ".md", ".rst", ".py", ".json", ".yaml", ".yml", ".ini",
    ".cfg", ".conf", ".toml", ".sh", ".xml", ".html", ".htm", ".js", ".ts",
    ".c", ".h", ".cpp", ".hpp", ".java", ".go", ".rs", ".sql", ".tex", ".env.example",
}
TABLE_EXT = {".csv", ".tsv"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".bmp", ".gif", ".tif", ".tiff", ".webp"}
PDF_EXT = {".pdf"}
DOCX_EXT = {".docx"}
XLSX_EXT = {".xlsx", ".xlsm"}

MAX_FILE_BYTES = 80 * 1024 * 1024
MAX_PDF_PAGES = 300
MAX_TEXT_CHARS = 400_000
CHUNK_CHARS = 1200
CHUNK_OVERLAP = 200
TABLE_CHUNK_CHARS = 1500

OBSOLETE_RE = re.compile(
    r"\bwithdrawn\b|\bsuperseded\s+by\b|\bobsolete\b|\bdeprecated\b|\bdo not use\b|\bretired\b",
    re.I,
)

OcrFn = Callable[["object", str], str]  # (PIL.Image, hint) -> text


@dataclass
class Doc:
    path: str            # corpus-relative, forward slashes
    kind: str            # text / table / pdf / docx / xlsx / image
    text: str
    obsolete: bool = False
    note: str = ""


@dataclass
class Chunk:
    id: int
    path: str
    text: str


@dataclass
class IndexReport:
    indexed: list = field(default_factory=list)
    skipped: list = field(default_factory=list)   # (path, reason)


# --------------------------------------------------------------------------- walk

def walk(corpus: Path, report: IndexReport):
    """Yield (abs_path, rel_path) for every regular file, never raising."""
    def onerror(err):
        report.skipped.append((str(getattr(err, "filename", "?")), f"walk: {err!r}"))

    for root, dirs, files in os.walk(corpus, onerror=onerror, followlinks=False):
        dirs.sort()
        for name in sorted(files):
            p = Path(root) / name
            try:
                rel = p.relative_to(corpus).as_posix()
            except Exception:
                rel = name
            try:
                if p.is_symlink() or not p.is_file():
                    report.skipped.append((rel, "not a regular file"))
                    continue
            except Exception as e:
                report.skipped.append((rel, f"stat: {e!r}"))
                continue
            yield p, rel


# --------------------------------------------------------------------------- parsers

def _read_bytes(p: Path) -> bytes:
    size = p.stat().st_size
    if size > MAX_FILE_BYTES:
        raise ValueError(f"too large ({size} bytes)")
    with open(p, "rb") as fh:          # raises PermissionError when unreadable
        return fh.read()


def _decode(b: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "latin-1"):
        try:
            s = b.decode(enc)
            if enc == "utf-16" and not (b[:2] in (b"\xff\xfe", b"\xfe\xff")):
                continue
            return s
        except Exception:
            continue
    return b.decode("utf-8", errors="replace")


def _looks_binary(b: bytes) -> bool:
    if not b:
        return False
    sample = b[:4096]
    if b"\x00" in sample:
        return True
    ctrl = sum(1 for c in sample if c < 9 or (13 < c < 32))
    return ctrl / len(sample) > 0.05


def parse_text(b: bytes) -> str:
    if _looks_binary(b):
        raise ValueError("binary content in text file")
    return _decode(b)


def parse_table(b: bytes, ext: str) -> str:
    text = _decode(b)
    delim = "\t" if ext == ".tsv" else None
    if delim is None:
        try:
            delim = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|").delimiter
        except Exception:
            delim = ","
    rows = list(csv.reader(io.StringIO(text), delimiter=delim))
    return _rows_to_text(rows)


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _rows_to_text(rows, sheet: Optional[str] = None) -> str:
    """Tables become one line per row, as `header=value` pairs, so each row is
    self-describing (good for both retrieval and the LLM)."""
    rows = [[_cell(c) for c in r] for r in rows]
    rows = [r for r in rows if any(r)]
    if not rows:
        return ""
    lines = []
    if sheet:
        lines.append(f"## Sheet: {sheet}")
    header = rows[0]
    has_header = len(rows) > 1 and all(h and not _is_number(h) for h in header)
    if has_header:
        lines.append("columns: " + " | ".join(header))
        for r in rows[1:]:
            pairs = []
            for i, v in enumerate(r):
                if not v:
                    continue
                h = header[i] if i < len(header) and header[i] else f"col{i+1}"
                pairs.append(f"{h}={v}")
            lines.append("; ".join(pairs))
    else:
        for r in rows:
            lines.append(" | ".join(v for v in r if v))
    return "\n".join(lines)


def _is_number(s: str) -> bool:
    try:
        float(s.replace(",", ""))
        return True
    except Exception:
        return False


def parse_xlsx(p: Path) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    parts = []
    try:
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            t = _rows_to_text(rows, sheet=ws.title)
            if t:
                parts.append(t)
    finally:
        try:
            wb.close()
        except Exception:
            pass
    return "\n\n".join(parts)


def parse_docx(p: Path) -> str:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = docx.Document(str(p))
    out = []

    def table_text(t):
        rows = []
        for row in t.rows:
            cells = []
            prev = None
            for c in row.cells:            # merged cells repeat; drop dupes
                if prev is not None and c._tc is prev:
                    continue
                prev = c._tc
                cells.append(c.text.strip())
            rows.append(cells)
        return _rows_to_text(rows)

    body = d.element.body
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            t = Paragraph(child, d).text.strip()
            if t:
                out.append(t)
        elif tag == "tbl":
            t = table_text(Table(child, d))
            if t:
                out.append(t)
    for s in d.sections:                    # headers / footers
        for part in (s.header, s.footer):
            try:
                for para in part.paragraphs:
                    if para.text.strip():
                        out.append(para.text.strip())
            except Exception:
                pass
    return "\n".join(out)


class EncryptedFile(Exception):
    pass


def parse_pdf(p: Path, ocr: Optional[OcrFn], ocr_allowed: Callable[[], bool]) -> str:
    import pypdf
    r = pypdf.PdfReader(str(p))
    if r.is_encrypted:
        # Deliberately NOT trying an empty password: a file we were not given
        # the password for is not a source, even if it happens to open.
        raise EncryptedFile("encrypted PDF")
    pages = []
    for i, page in enumerate(r.pages[:MAX_PDF_PAGES]):
        try:
            t = page.extract_text() or ""
        except Exception:
            t = ""
        if len(t.strip()) < 25 and ocr is not None and ocr_allowed():
            t2 = _ocr_pdf_page(p, i, ocr)
            if t2:
                t = t2
        if t.strip():
            pages.append(f"[page {i+1}]\n{t.strip()}")
    return "\n\n".join(pages)


def _ocr_pdf_page(p: Path, index: int, ocr: OcrFn) -> str:
    try:
        import pypdfium2 as pdfium
    except Exception:
        return ""
    try:
        doc = pdfium.PdfDocument(str(p))
        page = doc[index]
        img = page.render(scale=2.0).to_pil()
        return ocr(img, "scanned document page")
    except Exception:
        return ""


def parse_image(p: Path, ocr: Optional[OcrFn]) -> str:
    if ocr is None:
        raise ValueError("no OCR model available")
    from PIL import Image
    with Image.open(p) as im:
        im.load()
        img = im.convert("RGB")
    return ocr(img, "image")


# --------------------------------------------------------------------------- driver

def parse_file(p: Path, rel: str, ocr: Optional[OcrFn], ocr_allowed: Callable[[], bool]) -> Doc:
    ext = p.suffix.lower()
    name = p.name.lower()
    if ext in TEXT_EXT or name in {"readme", "makefile", "dockerfile"}:
        kind, text = "text", parse_text(_read_bytes(p))
    elif ext in TABLE_EXT:
        kind, text = "table", parse_table(_read_bytes(p), ext)
    elif ext in PDF_EXT:
        _read_bytes(p)  # permission check before handing to pypdf
        kind, text = "pdf", parse_pdf(p, ocr, ocr_allowed)
    elif ext in DOCX_EXT:
        _read_bytes(p)
        kind, text = "docx", parse_docx(p)
    elif ext in XLSX_EXT:
        _read_bytes(p)
        kind, text = "xlsx", parse_xlsx(p)
    elif ext in IMAGE_EXT:
        _read_bytes(p)
        if not ocr_allowed():
            raise ValueError("OCR time budget exhausted")
        kind, text = "image", parse_image(p, ocr)
    else:
        raise ValueError(f"unknown file type {ext or '(none)'}")
    text = text[:MAX_TEXT_CHARS]
    head = "\n".join([l for l in text.splitlines() if l.strip()][:4])
    obsolete = bool(OBSOLETE_RE.search(p.stem.replace("_", " "))) or bool(
        re.search(r"\bwithdrawn\b|\bsuperseded\s+by\b|\bobsolete\b", head, re.I)
    )
    return Doc(path=rel, kind=kind, text=text, obsolete=obsolete)


def build_docs(corpus: Path, ocr: Optional[OcrFn] = None,
               ocr_deadline: Optional[float] = None, log=print):
    report = IndexReport()
    docs: list[Doc] = []

    def ocr_allowed():
        return ocr_deadline is None or time.time() < ocr_deadline

    for p, rel in walk(corpus, report):
        t0 = time.time()
        try:
            d = parse_file(p, rel, ocr, ocr_allowed)
            if not d.text.strip():
                report.skipped.append((rel, "no extractable text"))
                continue
            docs.append(d)
            report.indexed.append(rel)
            log(f"[index] ok   {rel} ({d.kind}, {len(d.text)} chars, {time.time()-t0:.1f}s)"
                + ("  [OBSOLETE]" if d.obsolete else ""))
        except BaseException as e:  # noqa: BLE001 - never let one file stop the walk
            if isinstance(e, (KeyboardInterrupt, SystemExit)):
                raise
            report.skipped.append((rel, f"{type(e).__name__}: {e}"))
            log(f"[index] skip {rel}: {type(e).__name__}: {e}")
    return docs, report


def chunk_doc(d: Doc, start_id: int) -> list[Chunk]:
    chunks = []
    lines = d.text.splitlines()
    limit = TABLE_CHUNK_CHARS if d.kind in ("table", "xlsx") else CHUNK_CHARS
    if len(d.text) <= limit:
        return [Chunk(start_id, d.path, d.text)]
    # tables: keep the header line(s) on every chunk
    header = []
    if d.kind in ("table", "xlsx"):
        header = [l for l in lines[:2] if l.startswith(("columns:", "## Sheet"))]
    buf, size = list(header), sum(len(h) + 1 for h in header)
    i = 0
    cid = start_id
    while i < len(lines):
        line = lines[i]
        if line.startswith("## Sheet") and d.kind == "xlsx":
            header = [line]
        if line.startswith("columns:") and d.kind in ("table", "xlsx"):
            header = header[:1] + [line] if header and header[0].startswith("## Sheet") else [line]
        if size + len(line) + 1 > limit and len(buf) > len(header):
            chunks.append(Chunk(cid, d.path, "\n".join(buf)))
            cid += 1
            if d.kind in ("table", "xlsx"):
                buf, size = list(header), sum(len(h) + 1 for h in header)
            else:  # overlap for prose
                back, keep = 0, []
                for l in reversed(buf):
                    if back + len(l) > CHUNK_OVERLAP:
                        break
                    keep.insert(0, l)
                    back += len(l) + 1
                buf, size = keep, back
        buf.append(line[:limit])
        size += len(line[:limit]) + 1
        i += 1
    if buf and len(buf) > len(header):
        chunks.append(Chunk(cid, d.path, "\n".join(buf)))
    return chunks


def to_jsonable(docs, chunks, report):
    return {
        "docs": [asdict(d) for d in docs],
        "chunks": [asdict(c) for c in chunks],
        "report": {"indexed": report.indexed, "skipped": report.skipped},
    }
